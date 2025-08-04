from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from .models import Tender
from .serializers import TenderSerializer
from bids.models import Bid
from ai.compliance.engine import run_compliance_check, rank_bids_for_tender
from bids.serializers import BidSerializer
from notifications.utils import notify_user
from django.utils import timezone
from datetime import timedelta
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from django.http import HttpResponse
from io import BytesIO
from django.urls import reverse
from rest_framework.decorators import api_view
from rest_framework import filters


class TenderViewSet(viewsets.ModelViewSet):
    queryset = Tender.objects.all().order_by('-created_at')
    serializer_class = TenderSerializer
    permission_classes = [permissions.IsAuthenticated]
    filter_backends = [filters.SearchFilter]
    search_fields = ['title', 'description', 'id', 'created_by__username']

    def perform_create(self, serializer):
        tender = serializer.save(created_by=self.request.user)
        # Notify the creator    
        notify_user(
            recipient=self.request.user,
            message=f"Tender '{tender.title}' was created successfully.",
            data={"type": "tender_created", "tender_id": tender.id}
        )
    
    @action(detail=True, methods=["post"], url_path="process_bids")
    def process_bids(self, request, pk=None):
        tender = self.get_object()
        bids = Bid.objects.filter(tender=tender)

        if not bids.exists():
            return Response({"detail": "No bids submitted yet."}, status=400)

        for bid in bids:
            run_compliance_check(bid)

        rank_bids_for_tender(tender.id)

        return Response({"detail": "Bids processed and ranked successfully."})
    
    @action(detail=True, methods=['get'])
    def bids(self, request, pk=None):
        tender = self.get_object()
        bids = Bid.objects.filter(tender=tender)
        page = self.paginate_queryset(bids)
        if page is not None:
            serializer = BidSerializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = BidSerializer(bids, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path="summary")
    def summary(self, request):
        now = timezone.now()
        next_month = now + timedelta(days=30)

        summary = {
            "total": Tender.objects.count(),
            "draft": Tender.objects.filter(status="draft").count(),
            "published": Tender.objects.filter(status="published").count(),
            "closed": Tender.objects.filter(status="closed").count(),
            "due_next_month": Tender.objects.filter(deadline__lte=next_month, deadline__gte=now).count(),
            "report_url": request.build_absolute_uri(reverse("tender-summary-report"))
        }
        return Response(summary)
        
@api_view(['GET'])
def tender_summary_excel(request):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font
    from io import BytesIO
    from django.http import HttpResponse

    tenders = Tender.objects.select_related("created_by").all()
    now_ = timezone.now()
    next_month = now_ + timedelta(days=30)

    summary_data = {
        "Total Tenders": tenders.count(),
        "Draft": tenders.filter(status="draft").count(),
        "Published": tenders.filter(status="published").count(),
        "Closed": tenders.filter(status="closed").count(),
        "Due Next Month": tenders.filter(deadline__lte=next_month, deadline__gte=now_).count(),
    }

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Tender Summary"
    summary_ws.append(["Metric", "Value"])
    for k, v in summary_data.items():
        summary_ws.append([k, v])

    for col in summary_ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        summary_ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4

    tenders_ws = wb.create_sheet("Tenders")
    tenders_ws.append(["Title", "Status", "Deadline", "Created By", "Created At"])
    for tender in tenders:
        tenders_ws.append([
            tender.title,
            tender.status,
            tender.deadline.strftime('%Y-%m-%d %H:%M'),
            tender.created_by.username if tender.created_by else "N/A",
            tender.created_at.strftime('%Y-%m-%d %H:%M'),
        ])

    for col in tenders_ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        tenders_ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="tender_summary.xlsx"'
    return response